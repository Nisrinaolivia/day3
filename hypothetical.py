import streamlit as st
from openai import OpenAI
from dotenv import load_dotenv
from docx import Document
import io

load_dotenv()
client = OpenAI()

st.title("Contract Formation - Hypothetical Answering Tool")

if "messages" not in st.session_state:
    st.session_state.messages = []

if "scenario_submitted" not in st.session_state:
    st.session_state.scenario_submitted = False

# ================================================================
# INSTRUCTION 1: The ability to accept a hypothetical scenario
# involving contract formation.
# ================================================================
st.subheader("1. Accept a hypothetical scenario")
scenario = st.text_area(
    "Describe the facts of your contract law scenario:",
    height=200,
    placeholder="e.g. On 1 March, Alice emailed Bob offering to sell her car for $5,000...",
)

if st.button("Analyse scenario") and scenario:
    # ============================================================
    # INSTRUCTION 2: The ability to advise on whether a contract
    # has been formed with reference to each of the elements of
    # contract formation.
    # ============================================================
    system_prompt = """You are a contract law teaching assistant. Analyse the hypothetical scenario
    for whether a contract has been formed. Structure your answer under these headings:

    1. Offer
    2. Acceptance
    3. Consideration
    4. Intention to create legal relations
    5. Certainty of terms
    6. Conclusion - has a contract been formed?

    For each element, state whether it is satisfied, referencing the specific facts given.
    Be clear this is general educational analysis, not legal advice for a real situation."""

    with st.spinner("Analysing..."):
        response = client.responses.create(
            model="gpt-4o",
            input=f"{system_prompt}\n\nScenario: {scenario}",
        )

    st.session_state.messages = [
        {"role": "user", "content": f"Scenario: {scenario}"},
        {"role": "assistant", "content": response.output_text},
    ]
    st.session_state.scenario_submitted = True

if st.session_state.scenario_submitted:
    # ============================================================
    # INSTRUCTION 3: The ability to provide a response to the
    # user on the screen.
    # ============================================================
    st.subheader("2. Advise on contract formation (elements)")
    st.subheader("3. Response shown on screen")
    for msg in st.session_state.messages:
        if msg["role"] == "assistant":
            st.write(msg["content"])

    # ============================================================
    # INSTRUCTION 4: The ability for the user to ask at least one
    # follow-up question to the initial response.
    # ============================================================
    st.subheader("4. Ask a follow-up question")
    follow_up = st.text_input("Your follow-up question:")

    if st.button("Ask") and follow_up:
        st.session_state.messages.append({"role": "user", "content": follow_up})

        with st.spinner("Thinking..."):
            conversation_text = "\n\n".join(
                f"{m['role'].upper()}: {m['content']}" for m in st.session_state.messages
            )
            follow_up_response = client.responses.create(
                model="gpt-4o",
                input=f"Continue this legal Q&A conversation, staying consistent with earlier analysis:\n\n{conversation_text}",
            )

        st.session_state.messages.append(
            {"role": "assistant", "content": follow_up_response.output_text}
        )
        st.rerun()

    if len(st.session_state.messages) > 2:
        st.subheader("Full conversation (including follow-ups)")
        for msg in st.session_state.messages:
            role_label = "You" if msg["role"] == "user" else "Assistant"
            st.write(f"**{role_label}:** {msg['content']}")

    # ============================================================
    # INSTRUCTION 5: The ability for the user to download a
    # .docx (Word) version of the response.
    # ============================================================
    st.subheader("5. Download response as .docx")

    def create_docx(messages):
        doc = Document()
        doc.add_heading("Contract Formation Analysis", level=1)
        for msg in messages:
            role_label = "Scenario / Question" if msg["role"] == "user" else "Analysis"
            doc.add_heading(role_label, level=2)
            doc.add_paragraph(msg["content"])

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    docx_file = create_docx(st.session_state.messages)

    st.download_button(
        label="Download response as Word document",
        data=docx_file,
        file_name="contract_analysis.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )